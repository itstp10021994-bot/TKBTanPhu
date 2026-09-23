"""
Lưu / tải dữ liệu của ứng dụng lên SharePoint.

Hỗ trợ 2 cách kết nối (cấu hình trong .streamlit/secrets.toml hoặc mục
Secrets của Streamlit Cloud — xem HUONG_DAN_SHAREPOINT.md):

  1. [sharepoint]      — gọi thẳng Microsoft Graph API bằng 1 App
                          Registration (Azure AD / Entra ID).
  2. [power_automate]  — gửi dữ liệu tới các flow Power Automate có trigger
                          "When a HTTP request is received"; flow tự ghi/đọc
                          file trong thư viện SharePoint.

Mỗi lần lưu tạo 2 file cùng tên trong thư mục SharePoint:
  - <tên>.json : bản sao lưu ĐẦY ĐỦ (dùng để tải lại vào ứng dụng).
  - <tên>.xlsx : bản Excel dễ đọc (mỗi bảng 1 sheet) cho người dùng mở
                 trực tiếp trên SharePoint.
"""

from __future__ import annotations

import base64
import io
import json
import re
import time
import unicodedata
from datetime import date, datetime
from urllib.parse import quote, urlparse

import pandas as pd
import requests

PHIEN_BAN_DINH_DANG = 1
TIMEOUT = 60

# Các bảng nhập liệu (DataFrame trong st.session_state) được sao lưu
BANG_DU_LIEU = {
    "departments": "To_chuyen_mon",
    "teachers": "Giao_vien",
    "classes": "Lop_hoc",
    "rooms": "Phong_dac_biet",
    "activities": "Mon_hoc_phan_cong",
    "grade_times": "Thoi_gian_bieu",
    "exam_students": "HS_du_thi",
    "exam_rooms": "Phong_thi",
    "exam_subjects": "Mon_thi",
}
# Các giá trị đơn / cấu hình
GIA_TRI_CAU_HINH = [
    "cfg_num_days", "cfg_periods_per_day", "cfg_morning_count", "cfg_max_seconds",
    "cfg_school_name", "free_depts",
    "ct_period_spread", "ct_room_capacity", "ct_order_group", "ct_dept_free_session", "ct_no_gap",
]


# Bảng nhập liệu -> tên SharePoint List mặc định (đổi được trong Secrets:
# [sharepoint.lists] classes = "TenListKhac")
LIST_MAC_DINH = {
    "departments": "ToChuyenMon",
    "teachers": "GiaoVien",
    "classes": "LopHoc",
    "rooms": "PhongDacBiet",
    "activities": "MonHocPhanCong",
    "grade_times": "GioHocTheoKhoi",
    "exam_students": "DanhSachHocSinhThi",
    "exam_rooms": "DanhSachPhongThi",
    "exam_subjects": "MonThi",
}
# Cột kiểu số của ứng dụng (đọc từ List về thì đổi sang số)
COT_SO = {"Khối", "Tiết", "Nhóm thứ tự", "Số tiết/tuần", "Số phòng cùng loại", "Sức chứa", "Số cột bàn"}


def khoa_cot(ten) -> str:
    """Khoá so khớp tên cột/list: bỏ phần trong ngoặc, bỏ dấu, chữ thường.
    VD 'Họ và tên (không bắt buộc)' -> 'hovaten', 'Số tiết/tuần' -> 'sotiettuan'."""
    ten = re.sub(r"\(.*?\)", "", str(ten or "")).replace("đ", "d").replace("Đ", "D")
    ten = unicodedata.normalize("NFKD", ten)
    return re.sub(r"[^a-z0-9]", "", "".join(c for c in ten if not unicodedata.combining(c)).lower())


def _gia_tri_trong(v) -> bool:
    if v is None:
        return True
    try:
        if pd.isna(v):
            return True
    except (TypeError, ValueError):
        pass
    return str(v).strip() == ""


def _chuoi_gon(v) -> str:
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


class LoiLuuTru(Exception):
    """Lỗi khi lưu/tải dữ liệu — thông điệp đã viết sẵn cho người dùng."""


# ---------------------------------------------------------------------
# Đóng gói / giải nén dữ liệu phiên làm việc
# ---------------------------------------------------------------------
def _df_sang_json(df: pd.DataFrame) -> dict:
    df = df.astype(object).where(pd.notna(df), None)
    return {"columns": [str(c) for c in df.columns], "data": df.values.tolist()}


def _json_sang_df(obj: dict) -> pd.DataFrame:
    return pd.DataFrame(obj.get("data", []), columns=obj.get("columns", []))


def _mac_dinh_json(o):
    if isinstance(o, (date, datetime)):
        return o.isoformat()
    if hasattr(o, "model_dump"):
        return o.model_dump()
    if hasattr(o, "item"):  # numpy scalar
        return o.item()
    return str(o)


def dong_goi(ss) -> dict:
    """Gom toàn bộ dữ liệu cần lưu từ st.session_state thành 1 dict JSON."""
    goi = {
        "phien_ban": PHIEN_BAN_DINH_DANG,
        "thoi_gian_luu": datetime.now().isoformat(timespec="seconds"),
        "bang": {k: _df_sang_json(ss[k]) for k in BANG_DU_LIEU if k in ss},
        "cau_hinh": {k: ss[k] for k in GIA_TRI_CAU_HINH if k in ss},
        "tuan_bat_dau": ss["tuan_bat_dau"].isoformat() if "tuan_bat_dau" in ss else None,
        "ket_qua": {},
    }
    kq = ss.get("result")
    if kq is not None and kq.status not in ("INFEASIBLE", "ERROR"):
        goi["ket_qua"]["tkb"] = {
            "result": kq.model_dump(),
            "classes": [c.model_dump() for c in ss.get("result_classes", [])],
            "config": ss["result_config"].model_dump() if ss.get("result_config") else None,
        }
    if ss.get("exam_results"):
        goi["ket_qua"]["phong_thi"] = ss["exam_results"]
        goi["ket_qua"]["hoc_sinh_thi"] = ss.get("exam_all_students", [])
    if ss.get("substitutions"):
        goi["ket_qua"]["day_thay"] = ss["substitutions"]
    # đi qua json 1 lần để chắc chắn mọi giá trị đều tuần tự hoá được
    return json.loads(json.dumps(goi, default=_mac_dinh_json, ensure_ascii=False))


def giai_nen(goi: dict, ss, TimetableResult, SchoolClass, ScheduleConfig) -> list[str]:
    """Nạp dict đã lưu vào st.session_state. Truyền các class pydantic vào
    để module này không phụ thuộc vòng vào models. Trả về các ghi chú."""
    if not isinstance(goi, dict) or "bang" not in goi:
        raise LoiLuuTru("File không đúng định dạng sao lưu của ứng dụng.")
    ghi_chu = []
    for k, obj in goi.get("bang", {}).items():
        if k in BANG_DU_LIEU:
            ss[k] = _json_sang_df(obj)
            ss.pop(f"editor_{k}", None)
    for k, v in goi.get("cau_hinh", {}).items():
        if k in GIA_TRI_CAU_HINH:
            ss[k] = v
    if goi.get("tuan_bat_dau"):
        ss["tuan_bat_dau"] = date.fromisoformat(goi["tuan_bat_dau"])

    kq = goi.get("ket_qua", {})
    ss["result"], ss["solutions_history"], ss["selected_solution_idx"] = None, [], 0
    if kq.get("tkb"):
        tkb = kq["tkb"]
        result = TimetableResult.model_validate(tkb["result"])
        classes = [SchoolClass.model_validate(c) for c in tkb.get("classes", [])]
        config = ScheduleConfig.model_validate(tkb["config"]) if tkb.get("config") else ScheduleConfig()
        ss["result"], ss["result_classes"], ss["result_config"] = result, classes, config
        ss["solutions_history"] = [{"result": result, "classes": classes, "config": config}]
        ghi_chu.append("Đã nạp kết quả thời khoá biểu.")
    ss["exam_results"] = kq.get("phong_thi", {})
    ss["exam_all_students"] = kq.get("hoc_sinh_thi", [])
    ss["substitutions"] = kq.get("day_thay", {})
    if ss["exam_results"]:
        ghi_chu.append(f"Đã nạp kết quả xếp phòng thi ({len(ss['exam_results'])} môn).")
    return ghi_chu


def sang_excel(goi: dict) -> bytes:
    """Bản Excel dễ đọc: mỗi bảng nhập liệu 1 sheet, cộng các sheet kết quả."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        pd.DataFrame(
            [{"Mục": "Thời gian lưu", "Giá trị": goi.get("thoi_gian_luu")},
             {"Mục": "Tuần bắt đầu", "Giá trị": goi.get("tuan_bat_dau")}]
            + [{"Mục": k, "Giá trị": json.dumps(v, ensure_ascii=False) if isinstance(v, list) else v}
               for k, v in goi.get("cau_hinh", {}).items()]
        ).to_excel(writer, index=False, sheet_name="Thong_tin")
        for k, obj in goi.get("bang", {}).items():
            _json_sang_df(obj).to_excel(writer, index=False, sheet_name=BANG_DU_LIEU.get(k, k)[:31])

        kq = goi.get("ket_qua", {})
        if kq.get("tkb"):
            ten_lop = {c["id"]: c["name"] for c in kq["tkb"].get("classes", [])}
            pd.DataFrame([
                {"Lớp": ", ".join(ten_lop.get(c, c) for c in l["class_ids"]), "Thứ": l["day"] + 1,
                 "Tiết": l["period"], "Hoạt động": l["activity_name"], "GV": ", ".join(l["teacher_ids"])}
                for l in kq["tkb"]["result"].get("lessons", [])
            ]).sort_values(["Lớp", "Thứ", "Tiết"]).to_excel(writer, index=False, sheet_name="KQ_Thoi_khoa_bieu")
        if kq.get("phong_thi"):
            pd.DataFrame([
                {"Môn thi": mon, "Ngày thi": e.get("ngay"), "Ca thi": e.get("ca"),
                 "Phòng thi": p["ten_phong"], "SBD": hs["sbd"], "Lớp": hs["lop"]}
                for mon, e in kq["phong_thi"].items() for p in e.get("rooms", []) for hs in p["hoc_sinh"]
            ]).to_excel(writer, index=False, sheet_name="KQ_Phong_thi")
    return buf.getvalue()


def mau_excel_tao_list(bang_df: dict[str, pd.DataFrame], ten_list: dict[str, str]) -> bytes:
    """File Excel để tạo SharePoint List bằng 'Danh sách mới → Từ Excel':
    mỗi List 1 sheet, dữ liệu nằm trong 1 Excel Table cùng tên với List."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    wb.remove(wb.active)
    for bang, df in bang_df.items():
        ten = ten_list[bang]
        ws = wb.create_sheet(ten[:31])
        cot = [str(c) for c in df.columns]
        ws.append(cot)
        du_lieu = df.astype(object).where(pd.notna(df), None).values.tolist()
        for dong in du_lieu or [[None] * len(cot)]:
            ws.append([v.item() if hasattr(v, "item") else v for v in dong])
        for j, c in enumerate(cot, start=1):
            ws.column_dimensions[get_column_letter(j)].width = min(max(len(c) + 4, 12), 40)
        bang_excel = Table(displayName=re.sub(r"[^A-Za-z0-9_]", "_", ten) or f"Bang_{bang}",
                           ref=f"A1:{get_column_letter(len(cot))}{ws.max_row}")
        bang_excel.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(bang_excel)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ten_file_an_toan(ten: str) -> str:
    """Bỏ các ký tự SharePoint không cho phép trong tên file."""
    ten = re.sub(r'[\\/:*?"<>|#%~&{}]+', "_", (ten or "").strip()).strip(". _")
    return ten[:100] or datetime.now().strftime("TKB_%Y%m%d_%H%M%S")


# ---------------------------------------------------------------------
# Đồng bộ trực tiếp với SharePoint Lists: mỗi bảng nhập liệu <-> 1 List,
# mỗi dòng của bảng = 1 item. Cột khớp theo TÊN HIỂN THỊ (không phân biệt
# dấu, hoa/thường, bỏ phần trong ngoặc) — VD cột list "So tiet tuan" khớp
# cột "Số tiết/tuần" của ứng dụng. Lớp con (Graph / Power Automate) chỉ cần
# cài 4 hàm: _tim_list, _cot_list, _doc_items, _thay_items.
# ---------------------------------------------------------------------
class DongBoList:
    cfg: dict
    co_dong_bo_list = False
    co_luu_file = False

    def ten_list(self, bang: str) -> str:
        tuy_chinh = self.cfg.get("lists") or {}
        return str(tuy_chinh.get(bang) or LIST_MAC_DINH[bang])

    def _ghep_cot(self, ref, cot_app: list[str]) -> tuple[dict, list[str], dict | None]:
        """-> ({cột app: cột list}, [cột app không có trong list], cột Title)."""
        cot_list = self._cot_list(ref)
        ghep, thieu = {}, []
        for c in cot_app:
            k = khoa_cot(c)
            khop = next((x for x in cot_list
                         if khoa_cot(x["displayName"]) == k or khoa_cot(x["name"]) == k), None)
            if khop:
                ghep[c] = khop
            else:
                thieu.append(c)
        title = next((x for x in cot_list if x["name"] == "Title"), None)
        # Cột đầu tiên của bảng (Tên tổ / Tên giáo viên / Tên phòng...) không có trong
        # List -> dùng cột Tiêu đề (Title) mặc định, vì người dùng thường nhập tên ở đó.
        if (title is not None and cot_app and cot_app[0] in thieu
                and all(x["name"] != "Title" for x in ghep.values())):
            ghep = {cot_app[0]: title, **ghep}
            thieu.remove(cot_app[0])
        return ghep, thieu, title

    def _can_list(self, bang: str):
        ten = self.ten_list(bang)
        ref = self._tim_list(ten)
        if ref is None:
            raise LoiLuuTru(f"Không tìm thấy List '{ten}'.")
        return ten, ref

    def kiem_tra_list(self, bang_cot: dict[str, list[str]]) -> list[dict]:
        """Kiểm tra từng bảng: List có tồn tại không, cột nào khớp / thiếu."""
        ket_qua = []
        for bang, cot_app in bang_cot.items():
            ten = self.ten_list(bang)
            ref = self._tim_list(ten)
            if ref is None:
                ket_qua.append({"bang": bang, "list": ten, "co_list": False, "khop": [], "thieu": cot_app})
                continue
            ghep, thieu, _ = self._ghep_cot(ref, cot_app)
            ket_qua.append({"bang": bang, "list": ten, "co_list": True,
                            "khop": [f"{c} → {x['displayName']}" for c, x in ghep.items()],
                            "thieu": thieu})
        return ket_qua

    def ghi_list(self, bang: str, df: pd.DataFrame) -> tuple[int, list[str]]:
        """Thay TOÀN BỘ item của List bằng các dòng của df. -> (số dòng, cảnh báo)."""
        ten, ref = self._can_list(bang)
        ghep, thieu, title = self._ghep_cot(ref, [str(c) for c in df.columns])
        if not ghep:
            raise LoiLuuTru(f"List '{ten}' không có cột nào trùng tên với bảng — kiểm tra lại tên cột.")
        canh_bao = [f"List '{ten}' thiếu cột {thieu} — dữ liệu các cột này KHÔNG được lưu."] if thieu else []

        items = []
        for _, row in df.iterrows():
            fields = {}
            for c, cot in ghep.items():
                v = row[c]
                if _gia_tri_trong(v):
                    continue
                if cot["kieu"] == "so":
                    so = pd.to_numeric(v, errors="coerce")
                    if pd.isna(so):
                        raise LoiLuuTru(f"List '{ten}', cột '{c}': '{v}' không phải là số.")
                    fields[cot["name"]] = int(so) if float(so).is_integer() else float(so)
                elif cot["kieu"] == "bool":
                    fields[cot["name"]] = str(v).strip().lower() in ("1", "true", "x", "có", "co", "yes")
                else:
                    fields[cot["name"]] = _chuoi_gon(v)
            if not fields:
                continue  # dòng trống
            if title is not None and "Title" not in fields:
                # cột Tiêu đề mặc định của List: điền giá trị cột đầu tiên cho dễ nhìn
                fields["Title"] = _chuoi_gon(next(iter(fields.values())))[:255]
            items.append(fields)

        ids_cu = [i for i, _ in self._doc_items(ref)]
        self._thay_items(ref, ids_cu, items)
        return len(items), canh_bao

    def doc_list(self, bang: str, cot_app: list[str]) -> tuple[pd.DataFrame, list[str]]:
        """Đọc List -> DataFrame đúng các cột của ứng dụng. -> (df, cảnh báo)."""
        ten, ref = self._can_list(bang)
        ghep, thieu, _ = self._ghep_cot(ref, cot_app)
        rows = []
        for _, f in self._doc_items(ref):
            rows.append({c: ("" if _gia_tri_trong(f.get(ghep[c]["name"])) else f.get(ghep[c]["name"]))
                         if c in ghep else "" for c in cot_app})
        df = pd.DataFrame(rows, columns=cot_app)
        for c in cot_app:
            if c in COT_SO:
                so = pd.to_numeric(df[c], errors="coerce")
                df[c] = so.astype(int) if len(so) and so.notna().all() and (so % 1 == 0).all() else so
        canh_bao = [f"List '{ten}' thiếu cột {thieu} — để trống khi tải về."] if thieu else []
        return df, canh_bao


# ---------------------------------------------------------------------
# Cách 1: Microsoft Graph API (gọi thẳng SharePoint)
# ---------------------------------------------------------------------
class GraphSharePoint(DongBoList):
    ten_hien_thi = "SharePoint (Microsoft Graph)"

    def __init__(self, cfg: dict):
        thieu = [k for k in ("tenant_id", "client_id", "client_secret", "site_url") if not cfg.get(k)]
        if thieu:
            raise LoiLuuTru(f"Mục [sharepoint] trong Secrets còn thiếu: {', '.join(thieu)}.")
        self.cfg = cfg
        self.thu_muc = str(cfg.get("folder") or "TKB_TanPhu").strip("/")
        self._token = None
        self._drive_id = None
        self._site_id = None
        self._lists = None
        self._cot_cache: dict[str, list[dict]] = {}

    def _lay_token(self) -> str:
        if self._token:
            return self._token
        r = requests.post(
            f"https://login.microsoftonline.com/{self.cfg['tenant_id']}/oauth2/v2.0/token",
            data={
                "grant_type": "client_credentials",
                "client_id": self.cfg["client_id"],
                "client_secret": self.cfg["client_secret"],
                "scope": "https://graph.microsoft.com/.default",
            },
            timeout=TIMEOUT,
        )
        if r.status_code != 200:
            raise LoiLuuTru(
                "Đăng nhập Microsoft thất bại — kiểm tra tenant_id / client_id / client_secret "
                f"(chi tiết: {r.text[:300]})"
            )
        self._token = r.json()["access_token"]
        return self._token

    def _goi(self, method: str, url: str, **kw) -> requests.Response:
        headers = kw.pop("headers", {})
        headers["Authorization"] = f"Bearer {self._lay_token()}"
        if not url.startswith("https://"):
            url = f"https://graph.microsoft.com/v1.0{url}"
        r = requests.request(method, url, headers=headers, timeout=TIMEOUT, **kw)
        if r.status_code == 403:
            raise LoiLuuTru(
                "Ứng dụng chưa được cấp quyền ghi vào site SharePoint này (403). Xem bước cấp "
                "quyền Sites.ReadWrite.All / Sites.Selected trong HUONG_DAN_SHAREPOINT.md."
            )
        return r

    def _site(self) -> str:
        if self._site_id:
            return self._site_id
        u = urlparse(self.cfg["site_url"])
        duong_dan = u.path.rstrip("/")
        # link dán từ trình duyệt có thể kèm /Lists/..., /SitePages/... -> chỉ giữ phần site
        m = re.match(r"^(/(?:sites|teams|personal)/[^/]+)", duong_dan)
        if m:
            duong_dan = m.group(1)
        r = self._goi("GET", f"/sites/{u.hostname}:{duong_dan}" if duong_dan else f"/sites/{u.hostname}")
        if r.status_code != 200:
            raise LoiLuuTru(f"Không tìm thấy site SharePoint '{self.cfg['site_url']}' ({r.status_code}).")
        self._site_id = r.json()["id"]
        return self._site_id

    def _drive(self) -> str:
        if self._drive_id:
            return self._drive_id
        site_id = self._site()
        ten_thu_vien = self.cfg.get("library")
        if ten_thu_vien:
            r = self._goi("GET", f"/sites/{site_id}/drives")
            drives = [d for d in r.json().get("value", []) if d.get("name") == ten_thu_vien]
            if not drives:
                raise LoiLuuTru(f"Site không có thư viện tài liệu tên '{ten_thu_vien}'.")
            self._drive_id = drives[0]["id"]
        else:
            r = self._goi("GET", f"/sites/{site_id}/drive")
            self._drive_id = r.json()["id"]
        return self._drive_id

    def _duong_dan(self, ten_file: str) -> str:
        return quote(f"{self.thu_muc}/{ten_file}")

    def luu(self, ten: str, goi: dict) -> str:
        ten = ten_file_an_toan(ten)
        drive = self._drive()
        for duoi, noi_dung, loai in (
            ("json", json.dumps(goi, ensure_ascii=False, indent=1).encode("utf-8"), "application/json"),
            ("xlsx", sang_excel(goi),
             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ):
            r = self._goi("PUT", f"/drives/{drive}/root:/{self._duong_dan(f'{ten}.{duoi}')}:/content",
                          data=noi_dung, headers={"Content-Type": loai})
            if r.status_code not in (200, 201):
                raise LoiLuuTru(f"Lưu file {ten}.{duoi} thất bại ({r.status_code}): {r.text[:300]}")
        return f"{self.thu_muc}/{ten}.json"

    def danh_sach(self) -> list[dict]:
        r = self._goi("GET", f"/drives/{self._drive()}/root:/{quote(self.thu_muc)}:/children"
                             "?$select=name,lastModifiedDateTime&$top=200")
        if r.status_code == 404:
            return []  # thư mục chưa có -> chưa lưu lần nào
        if r.status_code != 200:
            raise LoiLuuTru(f"Không đọc được thư mục '{self.thu_muc}' ({r.status_code}).")
        return sorted(
            [{"ten": f["name"][:-5], "sua_luc": f.get("lastModifiedDateTime", "")}
             for f in r.json().get("value", []) if f["name"].lower().endswith(".json")],
            key=lambda x: x["sua_luc"], reverse=True,
        )

    def tai(self, ten: str) -> dict:
        r = self._goi("GET", f"/drives/{self._drive()}/root:/{self._duong_dan(f'{ten}.json')}:/content")
        if r.status_code != 200:
            raise LoiLuuTru(f"Không tải được '{ten}.json' ({r.status_code}).")
        return json.loads(r.content.decode("utf-8-sig"))


    # ---- Nguồn dữ liệu List cho DongBoList (qua Graph API) ----
    co_dong_bo_list = True
    co_luu_file = True

    def _tim_list(self, ten: str) -> str | None:
        if self._lists is None:
            self._lists, url = [], f"/sites/{self._site()}/lists?$select=id,displayName,name&$top=200"
            while url:
                r = self._goi("GET", url)
                if r.status_code != 200:
                    raise LoiLuuTru(f"Không đọc được danh sách List của site ({r.status_code}).")
                self._lists += r.json().get("value", [])
                url = r.json().get("@odata.nextLink")
        k = khoa_cot(ten)
        l = next((l for l in self._lists
                  if khoa_cot(l.get("displayName")) == k or khoa_cot(l.get("name")) == k), None)
        return l["id"] if l else None

    def _cot_list(self, list_id: str) -> list[dict]:
        if list_id not in self._cot_cache:
            r = self._goi("GET", f"/sites/{self._site()}/lists/{list_id}/columns")
            if r.status_code != 200:
                raise LoiLuuTru(f"Không đọc được cột của List ({r.status_code}).")
            self._cot_cache[list_id] = [
                {"name": c["name"], "displayName": c.get("displayName", c["name"]),
                 "kieu": "so" if "number" in c else "bool" if "boolean" in c else "text"}
                for c in r.json().get("value", [])
                if c.get("name") == "Title" or not (c.get("readOnly") or c.get("hidden"))
            ]
        return self._cot_cache[list_id]

    def _doc_items(self, list_id: str) -> list[tuple[str, dict]]:
        items, url = [], f"/sites/{self._site()}/lists/{list_id}/items?$expand=fields&$top=500"
        while url:
            r = self._goi("GET", url)
            if r.status_code != 200:
                raise LoiLuuTru(f"Không đọc được dữ liệu List ({r.status_code}): {r.text[:200]}")
            items += r.json().get("value", [])
            url = r.json().get("@odata.nextLink")
        items.sort(key=lambda it: int(it["id"]) if str(it.get("id", "")).isdigit() else 0)
        return [(str(it["id"]), it.get("fields", {})) for it in items]

    def _thay_items(self, list_id: str, ids_cu: list[str], items_moi: list[dict]):
        duong_dan = f"/sites/{self._site()}/lists/{list_id}/items"
        self._batch([{"method": "DELETE", "url": f"{duong_dan}/{i}"} for i in ids_cu], tuan_tu=False)
        self._batch([{"method": "POST", "url": duong_dan, "headers": {"Content-Type": "application/json"},
                      "body": {"fields": f}} for f in items_moi], tuan_tu=True)

    def _batch(self, yeu_cau: list[dict], tuan_tu: bool):
        """Gửi nhiều request qua /$batch (20 request/lần). tuan_tu=True: chạy
        lần lượt (dependsOn) để giữ đúng thứ tự dòng khi tạo item."""
        for dau in range(0, len(yeu_cau), 20):
            nhom = yeu_cau[dau:dau + 20]
            for lan in range(5):
                reqs = []
                for i, rq in enumerate(nhom):
                    rq = {**rq, "id": str(i + 1)}
                    if tuan_tu and i > 0:
                        rq["dependsOn"] = [str(i)]
                    reqs.append(rq)
                r = self._goi("POST", "/$batch", json={"requests": reqs})
                if r.status_code != 200:
                    raise LoiLuuTru(f"Ghi List thất bại ({r.status_code}): {r.text[:300]}")
                tra_ve = {x["id"]: x for x in r.json().get("responses", [])}
                loi = [(i, tra_ve.get(str(i + 1), {})) for i in range(len(nhom))
                       if tra_ve.get(str(i + 1), {}).get("status", 500) >= 300]
                if not loi:
                    break
                if all(x.get("status") in (424, 429, 503, 504) for _, x in loi) and lan < 4:
                    cho = max([int((x.get("headers") or {}).get("Retry-After", 2)) for _, x in loi] + [2])
                    time.sleep(min(cho, 30))
                    nhom = [nhom[i] for i, _ in loi]  # chỉ gửi lại các request lỗi
                    continue
                i, x = next(((i, x) for i, x in loi if x.get("status") != 424), loi[0])
                chi_tiet = ((x.get("body") or {}).get("error") or {}).get("message", "")
                raise LoiLuuTru(f"SharePoint từ chối ghi dòng {dau + i + 1} ({x.get('status')}): {chi_tiet}")


# ---------------------------------------------------------------------
# Cách 2: Power Automate (flow HTTP trigger ghi/đọc SharePoint)
# ---------------------------------------------------------------------
def _giai_thich_loi_xac_thuc_flow(r: requests.Response, url: str) -> str:
    """Power Automate trả 401/403 khi chặn request ngay ở trigger (flow chưa
    chạy). Đọc mã lỗi để chỉ đúng nguyên nhân."""
    try:
        data = r.json()
    except ValueError:
        data = {}
    # Flow ĐÃ chạy, nhưng bước "Send an HTTP request to SharePoint" bị SharePoint
    # từ chối (Response chuyển tiếp mã lỗi) -> body có trường "source" = URL đã gọi.
    nguon = str(data.get("source", "")) if isinstance(data, dict) else ""
    if "sharepoint.com" in nguon:
        loi_flow = []
        if re.search(r"/Lists/[^/]+/", nguon) or ".aspx" in nguon:
            loi_flow.append(
                "ô 'Site Address' đang là link của 1 List — sửa thành địa chỉ SITE (phần trước "
                f"'/Lists/'): {nguon.split('/Lists/')[0]}"
            )
        if "concat(" in nguon or "triggerBody()" in nguon:
            loi_flow.append(
                "ô 'Uri' đang chứa biểu thức dạng CHỮ THƯỜNG — xoá đi, nhập lại bằng nút fx (Insert "
                "expression) hoặc ghép chữ với ô 'list' / 'duong_dan' từ Dynamic content"
            )
        if not loi_flow:
            loi_flow.append("kiểm tra Site Address / Uri của các bước SharePoint trong flow")
        return (f"Flow đã chạy nhưng SharePoint từ chối ({r.status_code}) ở bước 'Send an HTTP request "
                f"to SharePoint': " + "; ".join(loi_flow) + f". [URL flow đã gọi: {nguon[:250]}]")
    loi = data.get("error", {}) if isinstance(data, dict) else {}
    loi = loi if isinstance(loi, dict) else {}
    ma, thong_diep = str(loi.get("code", "")), str(loi.get("message", "") or r.text[:200])
    chi_tiet = f" [mã lỗi: {ma or r.status_code}{' — ' + thong_diep[:200] if thong_diep else ''}]"
    if "sig=" not in url:
        goi_y = ("URL trong Secrets bị THIẾU phần '&sig=...' ở cuối (copy chưa hết). Copy lại toàn bộ "
                 "HTTP URL của trigger bằng nút copy bên cạnh ô URL.")
    elif "&amp;" in url:
        goi_y = "URL trong Secrets chứa '&amp;' — thay tất cả '&amp;' bằng '&'."
    elif any(k in ma for k in ("DirectApiAuthorizationRequired", "MisMatchingOAuthClaims", "OAuth",
                               "Unauthorized", "AuthorizationFailed")) or "OAuth" in thong_diep:
        goi_y = ("Trigger đang yêu cầu đăng nhập Microsoft. Mở flow → bấm vào trigger 'When a HTTP "
                 "request is received' → mục 'Who can trigger the flow' chọn 'Anyone' → Save, rồi COPY "
                 "LẠI HTTP URL (URL có thể đổi sau khi lưu) và cập nhật Secrets.")
    else:
        goi_y = ("Kiểm tra: (1) mục 'Who can trigger the flow' của trigger là 'Anyone'; (2) URL trong "
                 "Secrets copy đủ, đúng flow; (3) flow đang Bật (On); (4) tài khoản có giấy phép Premium.")
    return f"Power Automate từ chối yêu cầu ({r.status_code}). {goi_y}{chi_tiet}"


class PowerAutomate(DongBoList):
    """save_url/list_url/load_url: lưu/tải FILE sao lưu (3 flow).
    sp_url: 1 flow duy nhất đọc/ghi SharePoint List bằng quyền của chính
    người tạo flow — dùng được với 'Danh sách của tôi', không cần admin."""
    ten_hien_thi = "SharePoint qua Power Automate"
    SO_XOA_MOI_LAN = 200
    SO_TAO_MOI_LAN = 40
    COT_BO_QUA = {"ContentType", "Attachments"}

    def __init__(self, cfg: dict):
        if not (cfg.get("save_url") or cfg.get("sp_url")):
            raise LoiLuuTru("Mục [power_automate] trong Secrets cần có save_url và/hoặc sp_url.")
        self.cfg = cfg
        self.co_luu_file = bool(cfg.get("save_url"))
        self.co_dong_bo_list = bool(cfg.get("sp_url"))
        self._cot_cache: dict[str, list[dict]] = {}

    def _post(self, khoa_url: str, body: dict, cho_phep_404: bool = False) -> requests.Response:
        url = self.cfg.get(khoa_url)
        if not url:
            raise LoiLuuTru(f"Chưa cấu hình {khoa_url} trong mục [power_automate].")
        try:
            r = requests.post(url, json=body, timeout=230)
        except requests.Timeout:
            raise LoiLuuTru("Flow Power Automate chạy quá lâu, không phản hồi — xem lịch sử chạy (Run history) của flow.")
        if cho_phep_404 and r.status_code == 404:
            return r
        if r.status_code in (502, 504):
            raise LoiLuuTru(
                f"Flow Power Automate không trả kết quả ({r.status_code}) — mở Run history của flow để "
                "xem bước nào lỗi (thường do sai tên List/cột hoặc thiếu bước Response)."
            )
        if r.status_code in (401, 403):
            raise LoiLuuTru(_giai_thich_loi_xac_thuc_flow(r, url))
        if r.status_code >= 300:
            raise LoiLuuTru(f"Flow Power Automate trả lỗi {r.status_code}: {r.text[:300]}")
        return r

    # ---- lưu / tải file sao lưu ----
    def luu(self, ten: str, goi: dict) -> str:
        ten = ten_file_an_toan(ten)
        self._post("save_url", {
            "ten": ten,
            "thoi_gian": goi.get("thoi_gian_luu"),
            "json_text": json.dumps(goi, ensure_ascii=False),
            "excel_base64": base64.b64encode(sang_excel(goi)).decode("ascii"),
        })
        return f"{ten}.json"

    def co_the_tai(self) -> bool:
        return bool(self.cfg.get("list_url") and self.cfg.get("load_url"))

    def danh_sach(self) -> list[dict]:
        data = self._post("list_url", {}).json()
        if isinstance(data, dict):
            data = data.get("value") or data.get("files") or []
        ket_qua = []
        for f in data:
            ten = f.get("ten") or f.get("name") or f.get("Name") or ""
            if ten.lower().endswith(".json"):
                ten = ten[:-5]
            elif "." in ten:
                continue  # bỏ qua file .xlsx
            if ten:
                ket_qua.append({"ten": ten, "sua_luc": f.get("sua_luc") or f.get("LastModified") or ""})
        return sorted(ket_qua, key=lambda x: x["sua_luc"], reverse=True)

    def tai(self, ten: str) -> dict:
        r = self._post("load_url", {"ten": ten})
        data = r.json()
        if isinstance(data, str):  # flow trả về nội dung file dạng chuỗi
            data = json.loads(data)
        if isinstance(data, dict) and "$content" in data:  # nội dung nhị phân base64 của SharePoint
            data = json.loads(base64.b64decode(data["$content"]).decode("utf-8-sig"))
        return data

    # ---- đồng bộ List qua flow sp_url (SharePoint REST, odata=nometadata) ----
    @staticmethod
    def _gia_tri(r: requests.Response) -> list:
        data = r.json()
        if isinstance(data, str):
            data = json.loads(data)
        if isinstance(data, dict) and isinstance(data.get("d"), dict):  # odata=verbose
            data = data["d"].get("results", data["d"])
        if isinstance(data, dict):
            data = data.get("value", data.get("results", []))
        return data or []

    def _doc(self, ten: str, duong_dan: str, cho_phep_404: bool = False):
        return self._post("sp_url", {"thao_tac": "doc", "list": ten.replace("'", "''"),
                                     "duong_dan": duong_dan, "xoa_ids": [], "items": []},
                          cho_phep_404=cho_phep_404)

    def _tim_list(self, ten: str) -> str | None:
        if ten in self._cot_cache:
            return ten
        r = self._doc(ten, "fields?$filter=Hidden%20eq%20false%20and%20ReadOnlyField%20eq%20false"
                           "&$select=Title,InternalName,TypeAsString", cho_phep_404=True)
        if r.status_code == 404:
            return None
        self._cot_cache[ten] = [
            {"name": f["InternalName"], "displayName": f.get("Title") or f["InternalName"],
             "kieu": "so" if f.get("TypeAsString") in ("Number", "Integer", "Currency")
             else "bool" if f.get("TypeAsString") == "Boolean" else "text"}
            for f in self._gia_tri(r)
            if f.get("InternalName") and f["InternalName"] not in self.COT_BO_QUA
            and not str(f["InternalName"]).startswith("_")
        ]
        return ten

    def _cot_list(self, ten: str) -> list[dict]:
        self._tim_list(ten)
        return self._cot_cache.get(ten, [])

    def _doc_items(self, ten: str) -> list[tuple[str, dict]]:
        items = self._gia_tri(self._doc(ten, "items?$top=5000"))
        items.sort(key=lambda it: int(it.get("Id") or it.get("ID") or 0))
        return [(str(it.get("Id") or it.get("ID")), it) for it in items]

    def _thay_items(self, ten: str, ids_cu: list[str], items_moi: list[dict]):
        ten_sp = ten.replace("'", "''")
        for dau in range(0, len(ids_cu), self.SO_XOA_MOI_LAN):
            self._post("sp_url", {"thao_tac": "ghi", "list": ten_sp, "duong_dan": "",
                                  "xoa_ids": [int(i) for i in ids_cu[dau:dau + self.SO_XOA_MOI_LAN]],
                                  "items": []})
        for dau in range(0, len(items_moi), self.SO_TAO_MOI_LAN):
            self._post("sp_url", {"thao_tac": "ghi", "list": ten_sp, "duong_dan": "", "xoa_ids": [],
                                  "items": items_moi[dau:dau + self.SO_TAO_MOI_LAN]})


def tao_ket_noi(secrets) -> list:
    """Đọc st.secrets, trả về danh sách các kết nối đã cấu hình (có thể rỗng)."""
    ket_noi = []
    for khoa, lop in (("sharepoint", GraphSharePoint), ("power_automate", PowerAutomate)):
        try:
            cfg = dict(secrets[khoa]) if khoa in secrets else None
        except Exception:
            cfg = None
        if cfg:
            ket_noi.append((khoa, cfg, lop))
    return ket_noi
